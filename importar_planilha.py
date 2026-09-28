"""Lê fichas de treino de uma planilha (.xlsx ou .csv) enviada pelo usuário.

Devolve a mesma estrutura de fichas_academia.FICHAS: uma lista de dicts com
"name", "description" e "exercises" (tuplas nome, grupo, séries, reps,
descanso em segundos, observação).

O cabeçalho é reconhecido pelo nome das colunas (sem diferenciar maiúsculas
ou acentos) e pode estar em qualquer uma das primeiras linhas, então tanto o
modelo do app quanto a ficha da academia (título no topo, colunas "Volume"
3X12, "Descanso entre séries" 1 min, "Observações") funcionam. Sem coluna
"Ficha", cada aba do Excel vira uma ficha com o nome da aba.
"""

import csv
import io
import re
import unicodedata
from datetime import time

MAX_LINHAS = 500
MAX_FICHAS = 30
LINHAS_BUSCA_CABECALHO = 15


class PlanilhaInvalida(ValueError):
    """Erro com mensagem pronta para mostrar ao usuário."""


def _normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.lower().split())


def _classificar(cabecalho):
    h = _normalizar(cabecalho)
    if not h:
        return None
    if "descri" in h:
        return "description"
    if "descanso" in h or "intervalo" in h:
        return "rest_between" if "exerc" in h else "rest"
    if "observ" in h or h in ("obs", "obs."):
        return "notes"
    if "grupo" in h or "musculo" in h:
        return "group"
    if "volume" in h:
        return "volume"
    if "serie" in h:
        return "sets"
    if "repet" in h or h == "reps":
        return "reps"
    if "exerc" in h:
        return "exercise"
    if "ficha" in h or "treino" in h:
        return "plan"
    if "carga" in h or "peso" in h:
        return "ignore"
    return None


_CAMPOS_TABELA = {"exercise", "volume", "sets", "reps", "rest", "notes", "group"}


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    return "" if texto in ("-", "—", "–", "?") else texto


def _segundos(valor, em_minutos=False):
    """'1 min' -> 60, '90s' -> 90, '1:30' -> 90, 60 -> 60."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, time):
        return valor.hour * 3600 + valor.minute * 60 + valor.second
    if isinstance(valor, (int, float)):
        return int(round(valor * 60 if em_minutos else valor))
    texto = _normalizar(valor).replace(",", ".")
    m = re.fullmatch(r"(\d+):(\d{1,2})", texto)
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    m = re.match(r"(\d+(?:\.\d+)?)\s*(min|m\b|s|seg)?", texto)
    if not m:
        return None
    numero = float(m.group(1))
    unidade = m.group(2) or ("min" if em_minutos else "s")
    return int(round(numero * 60 if unidade.startswith("m") else numero))


def _inteiro(valor):
    m = re.match(r"\d+", _texto(valor))
    return int(m.group()) if m else None


def _ler_tabela(linhas, nome_padrao):
    """Converte as linhas de uma aba/arquivo em fichas."""
    linhas = [list(l) for l in linhas]

    inicio, colunas = None, {}
    for i, linha in enumerate(linhas[:LINHAS_BUSCA_CABECALHO]):
        tipos = {j: _classificar(c) for j, c in enumerate(linha)}
        if len({t for t in tipos.values() if t in _CAMPOS_TABELA}) >= 2:
            inicio, cabecalho = i, linha
            colunas = {t: j for j, t in reversed(tipos.items()) if t}
            break
    if inicio is None:
        return []

    dados = linhas[inicio + 1 :]
    if "exercise" not in colunas:
        # Na ficha da academia a coluna do nome do exercício não tem título:
        # usa a primeira coluna sem título que tenha conteúdo.
        usadas = set(colunas.values())
        for j in range(len(cabecalho)):
            if j not in usadas and _classificar(cabecalho[j]) is None and any(
                j < len(l) and _texto(l[j]) for l in dados
            ):
                colunas["exercise"] = j
                break
    if "exercise" not in colunas:
        raise PlanilhaInvalida("Não encontrei a coluna com o nome dos exercícios.")

    descanso_em_min = "min" in _normalizar(
        cabecalho[colunas["rest"]] if "rest" in colunas else ""
    )

    def celula(linha, campo):
        j = colunas.get(campo)
        return linha[j] if j is not None and j < len(linha) else None

    fichas, atual = {}, nome_padrao
    for linha in dados:
        if _texto(celula(linha, "plan")):
            atual = _texto(celula(linha, "plan"))
        nome = _texto(celula(linha, "exercise"))
        if not nome:
            continue
        if not atual:
            raise PlanilhaInvalida(
                f'O exercício "{nome}" está sem ficha. Preencha a coluna "Ficha".'
            )

        series = _inteiro(celula(linha, "sets"))
        reps = _texto(celula(linha, "reps"))
        volume = _normalizar(celula(linha, "volume"))
        m = re.fullmatch(r"(\d+)\s*[x×*]\s*(.+)", volume)
        if m:
            series = series or int(m.group(1))
            reps = reps or m.group(2).strip()

        ficha = fichas.setdefault(
            atual, {"name": atual, "description": None, "exercises": []}
        )
        if not ficha["description"]:
            descricao = _texto(celula(linha, "description"))
            entre_exercicios = _texto(celula(linha, "rest_between"))
            if descricao:
                ficha["description"] = descricao[:500]
            elif entre_exercicios:
                ficha["description"] = (
                    f"Descanso de {entre_exercicios} entre exercícios."
                )
        ficha["exercises"].append(
            (
                nome[:120],
                _texto(celula(linha, "group"))[:60] or None,
                min(series or 3, 20),
                (reps or "10")[:30],
                _segundos(celula(linha, "rest"), descanso_em_min),
                _texto(celula(linha, "notes"))[:255] or None,
            )
        )
    return list(fichas.values())


def _linhas_csv(conteudo):
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            texto = conteudo.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t")
    except csv.Error:
        dialeto = csv.excel
    return list(csv.reader(io.StringIO(texto), dialeto))


def ler_planilha(nome_arquivo, conteudo):
    """Lê o arquivo enviado e devolve a lista de fichas encontradas."""
    extensao = nome_arquivo.lower().rsplit(".", 1)[-1] if "." in nome_arquivo else ""
    nome_base = nome_arquivo.rsplit(".", 1)[0].strip() or "Ficha importada"

    if extensao == "csv":
        abas = [(nome_base, _linhas_csv(conteudo))]
    elif extensao == "xlsx":
        from openpyxl import load_workbook

        try:
            livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception as erro:
            raise PlanilhaInvalida("Não consegui abrir o arquivo do Excel.") from erro
        abas = [
            (aba.title, aba.iter_rows(max_row=MAX_LINHAS, values_only=True))
            for aba in livro.worksheets
        ]
    else:
        raise PlanilhaInvalida("Envie um arquivo .xlsx (Excel) ou .csv.")

    fichas = []
    for nome_aba, linhas in abas:
        fichas.extend(_ler_tabela(list(linhas)[:MAX_LINHAS], nome_aba))

    if not fichas:
        raise PlanilhaInvalida(
            "Não encontrei nenhuma ficha. Confira se a planilha tem o cabeçalho "
            "com as colunas (Exercício, Séries, Repetições...)."
        )
    if len(fichas) > MAX_FICHAS:
        raise PlanilhaInvalida(f"A planilha tem mais de {MAX_FICHAS} fichas.")
    return fichas
